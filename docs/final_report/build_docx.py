"""Build paper.docx from paper.tex in the IEEE two-column conference layout.

Run after compiling paper.tex (it reads paper.aux for figure, table, equation
and citation numbers):  python3 build_docx.py
Needs pandoc (only to turn the LaTeX math into Word equations) and python-docx.
"""
import copy
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
FIGS = HERE / "figs"
FONT = "Times New Roman"
PAGE_W, PAGE_H = 8.5, 11.0
MARGIN_LR, MARGIN_T, MARGIN_B = 0.625, 0.75, 1.0
COL_GAP = 0.2
COL_W = (PAGE_W - 2 * MARGIN_LR - COL_GAP) / 2
TEXT_W = PAGE_W - 2 * MARGIN_LR
TIKZ_PNG = {"fig:approach": "diagram_method.png", "fig:block": "diagram_block.png",
            "fig:flow": "diagram_flowchart.png"}
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"

tex = (HERE / "paper.tex").read_text()
aux = (HERE / "paper.aux").read_text()
LAB = dict(re.findall(r"\\newlabel\{([^}]*)\}\{\{([^}]*)\}", aux))
BIB = dict(re.findall(r"\\bibcite\{([^}]*)\}\{([^}]*)\}", aux))
tex = "\n".join(l for l in tex.splitlines() if not l.lstrip().startswith("%"))


def roman(n):
    out = ""
    for v, r in [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]:
        while n >= v:
            out += r
            n -= v
    return out


# ---------------------------------------------------------------- math -> OMML
def collect_math():
    text = re.sub(r"\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}", "", tex[tex.index(r"\begin{abstract}"):], flags=re.S)
    inline = re.findall(r"(?<!\\)\$([^$]+)\$", text)
    display = [re.sub(r"\\label\{[^}]*\}", "", e).strip()
               for e in re.findall(r"\\begin\{equation\}(.*?)\\end\{equation\}", tex, re.S)]
    items = list(dict.fromkeys(inline + display))
    md = "\n\n".join("\\(" + " ".join(x.split()) + "\\)" for x in items)
    with tempfile.TemporaryDirectory() as td:
        src, out = Path(td) / "m.tex", Path(td) / "m.docx"
        src.write_text(md)
        subprocess.run(["pandoc", "-f", "latex", str(src), "-o", str(out)], check=True)
        paras = [p for p in Document(str(out)).element.body.iter(qn("w:p"))]
    omml = {}
    for x, p in zip(items, paras):
        oms = list(p.iter(f"{M}oMath"))
        assert oms, x
        omml[x] = oms[0]
    return omml


OMML = collect_math()


# ------------------------------------------------------------ inline LaTeX text
def clean(t):
    t = t.replace("~", "\u00a0").replace(r"\%", "%").replace(r"\&", "&").replace(r"\_", "_")
    t = t.replace("``", "\u201c").replace("''", "\u201d").replace("--", "\u2013")
    t = t.replace(r"\'a", "\u00e1").replace(r"\,", "\u2009").replace(r"\times", "\u00d7")
    t = t.replace("\\\\", " ")
    t = re.sub(r"\\(footnotesize|centering|hfill)\b", "", t)
    t = re.sub(r"\s+", " ", t)
    return t


def runs_of(t, fmt=None):
    """Yield (kind, value, fmt) pieces: kind is 'text' or 'math'."""
    fmt = dict(fmt or {})
    pat = re.compile(r"(?<!\\)\$([^$]+)\$|\\(emph|textit|textbf|texttt|textsuperscript)\{"
                     r"|\\cite\{([^}]*)\}|\\eqref\{([^}]*)\}|\\ref\{([^}]*)\}")
    i = 0
    while i < len(t):
        m = pat.search(t, i)
        if not m:
            yield ("text", clean(t[i:]), fmt)
            return
        if m.start() > i:
            yield ("text", clean(t[i:m.start()]), fmt)
        if m.group(1) is not None:
            yield ("math", m.group(1), fmt)
            i = m.end()
        elif m.group(2):
            depth, j = 1, m.end()
            while depth:
                depth += {"{": 1, "}": -1}.get(t[j], 0)
                j += 1
            inner = t[m.end():j - 1]
            f = dict(fmt)
            f[{"emph": "italic", "textit": "italic", "textbf": "bold", "texttt": "mono",
               "textsuperscript": "sup"}[m.group(2)]] = True
            yield from runs_of(inner, f)
            i = j
        elif m.group(3):
            yield ("text", ", ".join(f"[{BIB[k.strip()]}]" for k in m.group(3).split(",")), fmt)
            i = m.end()
        elif m.group(4):
            yield ("text", f"({LAB[m.group(4)]})", fmt)
            i = m.end()
        else:
            yield ("text", LAB[m.group(5)], fmt)
            i = m.end()


def set_font(run, size=None, **fmt):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), "Courier New" if fmt.get("mono") else FONT)
    if size:
        run.font.size = Pt(size)
    run.bold = fmt.get("bold") or None
    run.italic = fmt.get("italic") or None
    run.font.superscript = fmt.get("sup") or None
    run.font.small_caps = fmt.get("smallcaps") or None


def add_inline(par, t, size=10, base=None):
    for kind, val, fmt in runs_of(t, base):
        if kind == "text":
            if val:
                set_font(par.add_run(val), size, **fmt)
        else:
            par._p.append(copy.deepcopy(OMML[val]))


def para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=0.2, before=0, after=0, size=10):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.first_line_indent = Inches(indent) if indent else None
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.0
    p._size = size
    return p


# --------------------------------------------------------------------- document
doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = FONT, Pt(10)
st.element.get_or_add_rPr()
rf = st.element.rPr.find(qn("w:rFonts"))
if rf is None:
    rf = OxmlElement("w:rFonts")
    st.element.rPr.insert(0, rf)
for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rf.set(qn(a), FONT)

sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(PAGE_W), Inches(PAGE_H)
sec.left_margin = sec.right_margin = Inches(MARGIN_LR)
sec.top_margin, sec.bottom_margin = Inches(MARGIN_T), Inches(MARGIN_B)


def set_cols(section, n):
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(COL_GAP * 1440)))


# title
title = re.search(r"\\title\{(.*?)\}\n", tex).group(1)
p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 14)
set_font(p.add_run(title), 24)

# authors: three columns like the department sample, then the shared affiliation
names = re.findall(r"\\IEEEauthorblockN\{([^}]*)\}\s*\\IEEEauthorblockA\{([^}]*)\}", tex)
tbl = doc.add_table(rows=1, cols=len(names))
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, (n, a) in zip(tbl.rows[0].cells, names):
    cell.width = Inches(TEXT_W / len(names))
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(2)
    set_font(cp.add_run(n), 11)
    for line in a.split("\\\\"):
        q = cell.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(0)
        set_font(q.add_run(line.strip()), 10)
aff = tex[tex.rindex(r"\linebreakand"):tex.index(r"\maketitle")]
aff = re.search(r"\\IEEEauthorblockA\{(.*?)\}\s*\}", aff, re.S).group(1)
for k, line in enumerate(aff.split("\\\\")):
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, 10 if k == 0 else 0, 0)
    set_font(p.add_run(line.strip()), 10)
p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 6)

body = doc.add_section(WD_SECTION.CONTINUOUS)
set_cols(body, 2)

# abstract and index terms
abstract = tex[tex.index(r"\begin{abstract}") + 16:tex.index(r"\end{abstract}")].strip()
p = para(doc, indent=0.2, size=9)
set_font(p.add_run("Abstract\u2014"), 9, bold=True, italic=True)
add_inline(p, abstract, 9, {"bold": True})
kw = tex[tex.index(r"\begin{IEEEkeywords}") + 20:tex.index(r"\end{IEEEkeywords}")].strip()
p = para(doc, indent=0.2, after=6, size=9)
set_font(p.add_run("Index Terms\u2014"), 9, bold=True, italic=True)
add_inline(p, kw, 9, {"bold": True})

main = tex[tex.index(r"\end{IEEEkeywords}") + 18:tex.index(r"\end{document}")]
block = re.compile(r"\\section\*?\{([^}]*)\}|\\subsection\{([^}]*)\}|"
                   r"\\begin\{(enumerate|equation|figure\*?|table|thebibliography)\}(.*?)\\end\{\3\}", re.S)
sec_n = sub_n = 0
eq_n = 0


def heading(text, numbered=True):
    global sec_n, sub_n
    if numbered:
        sec_n += 1
        sub_n = 0
        text = f"{roman(sec_n)}. {text}"
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, 8, 4)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), 10, smallcaps=True)


def subheading(text):
    global sub_n
    sub_n += 1
    p = para(doc, WD_ALIGN_PARAGRAPH.LEFT, 0, 6, 3)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(f"{chr(64 + sub_n)}. {text}"), 10, italic=True)


def caption(text, label, is_table=False):
    num = LAB[label]
    if is_table:
        p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, 6, 0, 8)
        p.paragraph_format.keep_with_next = True
        set_font(p.add_run(f"TABLE {num}"), 8, smallcaps=True)
        p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 3, 8)
        p.paragraph_format.keep_with_next = True
        add_inline(p, text, 8, {"smallcaps": True})
    else:
        p = para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 3, 8, 8)
        set_font(p.add_run(f"Fig.\u00a0{num}.\u2002"), 8)
        add_inline(p, text, 8)


def figure(content, wide):
    label = re.search(r"\\label\{([^}]*)\}", content).group(1)
    cap = re.search(r"\\caption\{(.*)\}\s*\\label", content, re.S).group(1)
    files = re.findall(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]*)\}", content)
    if not files:
        files = [TIKZ_PNG[label]]
    width = TEXT_W - 0.1 if wide else COL_W - 0.05
    if wide:
        prev = doc.add_section(WD_SECTION.CONTINUOUS)
        set_cols(prev, 1)
    from PIL import Image
    sizes = [Image.open(FIGS / f).size for f in files]
    # same height for side-by-side images, total width filling the space
    h = 0.93 * (width - 0.25 * (len(files) - 1)) / sum(w / hh for w, hh in sizes)
    if len(files) == 1 and not wide and label != "fig:flow":
        h = min(h, 3.3)
    if label == "fig:flow":
        h = min(h, 4.2)
    if len(files) == 1:
        p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, 6, 0)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(FIGS / files[0]), height=Inches(h))
    else:  # side by side in a borderless table so they never wrap
        t = doc.add_table(rows=1, cols=len(files))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, f, (w, hh) in zip(t.rows[0].cells, files, sizes):
            cell.width = Inches(h * w / hh + 0.1)
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(0)
            cp.add_run().add_picture(str(FIGS / f), height=Inches(h))
    caption(cap, label)
    if wide:
        nxt = doc.add_section(WD_SECTION.CONTINUOUS)
        set_cols(nxt, 2)


def borders(tbl, header_rows):
    def edge(cell, side, sz):
        tcpr = cell._tc.get_or_add_tcPr()
        b = tcpr.find(qn("w:tcBorders"))
        if b is None:
            b = OxmlElement("w:tcBorders")
            tcpr.append(b)
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), "000000")
        b.append(e)
    rows = tbl.rows
    for c in rows[0].cells:
        edge(c, "top", 8)
    for c in rows[header_rows - 1].cells:
        edge(c, "bottom", 4)
    for c in rows[-1].cells:
        edge(c, "bottom", 8)


def table(content):
    label = re.search(r"\\label\{([^}]*)\}", content).group(1)
    cap = re.search(r"\\caption\{(.*?)\}\s*\\label", content, re.S).group(1)
    caption(cap, label, is_table=True)
    tab = content[content.index(r"\toprule") + 8:content.index(r"\bottomrule")]
    head, rest = tab.split(r"\midrule", 1)
    rest = rest.replace(r"\midrule", "")
    rows = [r for r in head.split("\\\\") if r.strip()]
    nhead = len(rows)
    rows += [r for r in rest.split("\\\\") if r.strip()]
    cells = [[c.strip() for c in r.split("&")] for r in rows]
    ncol = max(len(r) for r in cells)
    spec = re.search(r"\\begin\{tabular\}\{([^}]*)\}", content).group(1)
    aligns = [c for c in spec if c in "lrc"]
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in zip(t.rows, cells):
        for k, (cell, txt) in enumerate(zip(r.cells, row + [""] * (ncol - len(row)))):
            cp = cell.paragraphs[0]
            cp.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT,
                            "c": WD_ALIGN_PARAGRAPH.CENTER}[aligns[k] if k < len(aligns) else "l"]
            cp.paragraph_format.space_after = Pt(0)
            add_inline(cp, txt, 8)
    lens = [max(len(clean(re.sub(r"\\[a-z]+\{|\$|\}", "", row[k])) if k < len(row) else "") for row in cells) + 2
            for k in range(ncol)]
    tblpr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    tw = COL_W - 0.05
    for k, col in enumerate(t.columns):
        w = Inches(tw * lens[k] ** 0.6 / sum(x ** 0.6 for x in lens))
        col.width = w
        for c in col.cells:
            c.width = w
    borders(t, nhead)
    para(doc, indent=0, after=4)


def equation(content):
    global eq_n
    eq_n += 1
    key = re.sub(r"\\label\{[^}]*\}", "", content).strip()
    p = para(doc, WD_ALIGN_PARAGRAPH.LEFT, 0, 4, 4)
    ts = p.paragraph_format.tab_stops
    ts.add_tab_stop(Inches(COL_W / 2), WD_TAB_ALIGNMENT.CENTER)
    ts.add_tab_stop(Inches(COL_W - 0.02), WD_TAB_ALIGNMENT.RIGHT)
    set_font(p.add_run("\t"), 10)
    p._p.append(copy.deepcopy(OMML[key]))
    set_font(p.add_run(f"\t({eq_n})"), 10)


def text_paragraphs(chunk):
    for part in re.split(r"\n\s*\n", chunk):
        part = re.sub(r"\\FloatBarrier", "", part).strip()
        if part:
            p = para(doc)
            add_inline(p, part)


pos = 0
for m in block.finditer(main):
    text_paragraphs(main[pos:m.start()])
    pos = m.end()
    if m.group(1) is not None:
        star = main[m.start():m.start() + 9] == r"\section*"
        heading(m.group(1), numbered=not star)
    elif m.group(2) is not None:
        subheading(m.group(2))
    else:
        env, content = m.group(3), m.group(4)
        if env == "enumerate":
            for k, item in enumerate(re.split(r"\\item", content)[1:], 1):
                p = para(doc, indent=-0.2)
                p.paragraph_format.left_indent = Inches(0.4)
                set_font(p.add_run(f"{k})\t"), 10)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(0.4))
                add_inline(p, item.strip())
        elif env == "equation":
            equation(content)
        elif env.startswith("figure"):
            figure(content, env == "figure*")
        elif env == "table":
            table(content)
        else:  # thebibliography
            heading("References", numbered=False)
            for key, ref in re.findall(r"\\bibitem\{([^}]*)\}(.*?)(?=\\bibitem|$)", content, re.S):
                p = para(doc, WD_ALIGN_PARAGRAPH.LEFT, indent=-0.3, size=8)
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
                set_font(p.add_run(f"[{BIB[key]}]\t"), 8)
                add_inline(p, ref.strip(), 8)
text_paragraphs(main[pos:])

out = HERE / "paper.docx"
doc.save(str(out))
print("wrote", out)
